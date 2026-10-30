"""Build the Lab 08 (Pump Curve) SMath worksheet and its Excel plot workbook.

Usage: python3 lab08.py OUTPUT_DIR
"""
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
import subprocess, tempfile, shutil
import numpy as np
from smgen import Sheet, UNITS

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)
GPM = "'gal/'min"
S = Sheet()


def vec(vals, unit):
    return f"mat({','.join(str(v) for v in vals)},{len(vals)},1)" + (('*' + unit) if unit else '')


def vcall(fn, n):
    return f"mat({','.join(f'{fn}({i})' for i in range(1, n + 1))},{n},1)"


def vh(n):
    return 20 * n + 20


# ------------------------------------------------------------------ data from the data sheets
NM, NC = 6, 10
z_in = [47, 59.5625, 71, 88.5, 95, 107]
P_md = [2.72, 2.78, 2.82, 3.14, 3.28, 3.42]
Q_mag = [1.6, 1.5, 1.4, 1.15, 1.05, 0.7]
W_mag = [28.3, 27.9, 27.6, 27.2, 26.6, 26]
U_W_mag = ['0.002*28.3', '0.002*27.9', '0.002*27.6', '0.002*27.2', '0.1', '0.002*26']
pct = [100, 90, 80, 70, 60, 50, 40, 30, 20, 0]
P_in = [-3.75, -3.125, -2.5, -2.5, -2.5, -2.5, -2.5, -2.5, -2.5, 3.75]
P_dis = [7.8, 8.8, 9.8, 10.6, 11.4, 11.8, 12.4, 12.8, 13.0, 17.9]
Q_c = [45.3, 40.86, 36.53, 31.35, 26.38, 22.17, 17.58, 12.73, 8.01, 0]
# Lab 6 Part B system curve (from Lab06_PipeFriction_2.xlsx)
Q_sys = [3.5, 3.125, 2.75, 2.5, 2.125, 1.75, 1.375, 1.0]
h_sys = [26.7705, 16.9468, 14.1112, 12.3403, 8.1172, 5.4916, 2.7178, 0.9410]
U_h_sys = [0.6752, 0.5052, 0.4632, 0.4391, 0.1753, 0.1188, 0.0592, 0.0218]

# ------------------------------------------------------------------ header
for t in ('Grant Bellon', 'Lab [8]', 'Pump Curve and Efficiency Measurement'):
    S.text(t, bold=True, h=26, gap=12)
S.space(10)
S.text('Introduction', bold=True, h=26)
S.text('The objective of this experiment was to measure the pump curves and efficiencies of two centrifugal '
       'pumps. A magnetically-driven pump was run through six vertical pipes of increasing height, and its '
       'discharge pressure, flowrate, and electrical power draw were recorded so the head at the pump '
       'discharge and at each pipe zenith could be compared. A transparent centrifugal pump testbench was run '
       'at full speed while its discharge valve was closed in steps from fully open to fully closed, recording '
       'the intake and discharge pressures and flowrate. The heads and efficiencies of both pumps were '
       'calculated, their uncertainties were propagated with the Kline-McClintock method, and the pump curves '
       'were overlaid on the Lab 6 system curve to find the operating points.')
S.space(8)

# ------------------------------------------------------------------ given data
S.text('Given Data', bold=True, h=26)
S.text('Environmental conditions')
S.row([('def', "P.atm := 1004.7*'mbar", 140, 29), ('def', "U.P_atm := 0.1*'mbar", 140, 29),
       ('def', 'T.atm_degC := 25.8', 130, 29), ('def', 'U.T_atm_degC := 1', 130, 29)])
S.row([('def', "g := 9.81*'m/'s^2", 120, 50)])
S.text('Water density sample (the volume and mass uncertainties were not recorded; the graduated cylinder '
       'and scale resolutions are assumed)')
S.row([('def', "Vol.water := 7.6*'mL", 150, 29), ('def', "U.Vol_water := 0.1*'mL", 160, 29),
       ('def', "m.water := 7.18*'g", 140, 29), ('def', "U.m_water := 0.01*'g", 150, 29)])

S.text('Magnetic drive pump. Each vector lists vertical pipes 1 to 6. The zenith heights were measured with a '
       'tape measure (±1/16 in assumed). The wattmeter uncertainty is ±0.2% of the reading, except pipe 5, '
       'which was recorded as ±0.1 W.')
S.row([('def', "D.pipe := 21.15*'mm", 160, 29), ('def', "U.D_pipe := 0.01*'mm", 160, 29)])
S.row([('def', f"z.zenith := {vec(z_in, chr(39) + 'in')}", 250, vh(NM)),
       ('def', f"P.mag_discharge := {vec(P_md, chr(39) + 'psi')}", 270, vh(NM)),
       ('def', f"Q.magnetic := {vec(Q_mag, GPM)}", 260, vh(NM))], gap=10)
S.row([('def', "U.z_zenith := 1/16*'in", 250, 42), ('def', "U.P_mag_discharge := 0.1*'psi", 270, 42),
       ('def', "U.Q_magnetic := 0.05*" + GPM, 260, 42)], gap=16)
S.row([('def', f"W.magnetic := {vec(W_mag, chr(39) + 'W')}", 250, vh(NM)),
       ('def', f"U.W_magnetic := {vec(U_W_mag, chr(39) + 'W')}", 330, vh(NM))], gap=22)

S.text('Transparent centrifugal pump. Each vector lists the settings 100%, 90%, 80%, 70%, 60%, 50%, 40%, 30%, '
       '20%, and 0% (valve closed) of the maximum flowrate. The intake gauge reads in inches of mercury '
       '(negative = vacuum); the discharge gauge reads in psi.')
S.row([('def', "inHg := 3386.389*'Pa", 170, 29)])
S.row([('def', f"P.cent_intake := {vec(P_in, 'inHg')}", 260, vh(NC)),
       ('def', f"P.cent_discharge := {vec(P_dis, chr(39) + 'psi')}", 270, vh(NC)),
       ('def', f"Q.centrifugal := {vec(Q_c, GPM)}", 260, vh(NC))], gap=10)
S.row([('def', 'U.P_cent_intake := 0.625*inHg', 250, 30), ('def', "U.P_cent_discharge := 0.1*'psi", 270, 30),
       ('def', 'U.Q_centrifugal := 0.01*Q.centrifugal', 280, 30)], gap=16)
S.row([('def', "V.centrifugal := 237.9*'V", 180, 29), ('def', "U.V_centrifugal := 0.1*'V", 180, 29),
       ('def', "I.centrifugal := 4.1*'A", 170, 29), ('def', "U.I_centrifugal := 0.1*'A", 170, 29)], gap=22)

S.text('Lab 6 system curve (Part B of Lab06_PipeFriction_2.xlsx), tests 1 to 8 at 100% down to 30% of the '
       'maximum flowrate')
S.row([('def', f"Q.system := {vec(Q_sys, GPM)}", 260, vh(8)),
       ('def', f"h.L_system := {vec(h_sys, chr(39) + 'm')}", 260, vh(8)),
       ('def', f"U.h_L_system := {vec(U_h_sys, chr(39) + 'm')}", 270, vh(8))])

# ------------------------------------------------------------------ density
S.space(10)
S.text('Water Density', bold=True, h=26)
S.row([('defq', 'ρ.water := m.water/Vol.water', 330, 55, "'kg/'m^3")])
S.text('Kline-McClintock uncertainty of the density')
S.row([('defq', 'U.ρ_water := sqrt((U.m_water/Vol.water)^2+(m.water*U.Vol_water/Vol.water^2)^2)', 560, 75, "'kg/'m^3")])

# ------------------------------------------------------------------ magnetic drive pump
S.space(10)
S.text('Pump Curve of Magnetic Drive Pump', bold=True, h=26)
S.text('Specific energy and head at the pump discharge port for pipe k')
S.row([('fn', 'e.mag_discharge_k(k) := el(P.mag_discharge,k)/ρ.water', 380, 55),
       ('fn', 'h.mag_discharge_k(k) := e.mag_discharge_k(k)/g', 340, 55)])
S.row([('defq', 'h.mag_discharge := ' + vcall('h.mag_discharge_k', NM), 520, vh(NM), "'m")])
S.text('Pipe area and zenith velocity')
S.row([('defq', 'A.pipe := π*D.pipe^2/4', 260, 55, "'mm^2"),
       ('fn', 'V.zenith_k(k) := el(Q.magnetic,k)/A.pipe', 300, 55)])
S.row([('defq', 'V.zenith := ' + vcall('V.zenith_k', NM), 460, vh(NM), "'m/'s")])
S.text('Specific energy and head at the pipe zenith (P_zenith = 0, open to atmosphere)')
S.row([('fn', 'e.mag_zenith_k(k) := V.zenith_k(k)^2/2+g*el(z.zenith,k)', 400, 55),
       ('fn', 'h.mag_zenith_k(k) := e.mag_zenith_k(k)/g', 320, 55)])
S.row([('defq', 'h.mag_zenith := ' + vcall('h.mag_zenith_k', NM), 500, vh(NM), "'m")])
S.text('Efficiency, η = ṁ·e_zenith / Ẇ, with mass flowrate ṁ = ρ·Q')
S.row([('fn', 'm.dot_mag_k(k) := ρ.water*el(Q.magnetic,k)', 320, 35),
       ('fn', 'η.mag_k(k) := m.dot_mag_k(k)*e.mag_zenith_k(k)/el(W.magnetic,k)', 440, 55)])
S.row([('defq', 'η.magnetic := ' + vcall('η.mag_k', NM), 440, vh(NM))])

# ------------------------------------------------------------------ centrifugal pump
S.space(10)
S.text('Pump Curve of Transparent Centrifugal Pump', bold=True, h=26)
S.text('Specific energy and head added by the pump at setting k')
S.row([('fn', 'ΔP.cent_k(k) := el(P.cent_discharge,k)-el(P.cent_intake,k)', 440, 35)])
S.row([('fn', 'e.cent_k(k) := ΔP.cent_k(k)/ρ.water', 300, 55),
       ('fn', 'h.cent_k(k) := e.cent_k(k)/g', 260, 55)])
S.row([('defq', 'h.centrifugal := ' + vcall('h.cent_k', NC), 520, vh(NC), "'m")])
S.text('Electrical power draw and efficiency, η = ṁ·e / (I·V)')
S.row([('defq', 'W.dot_centrifugal := I.centrifugal*V.centrifugal', 380, 35, "'W")])
S.row([('fn', 'η.cent_k(k) := ρ.water*el(Q.centrifugal,k)*e.cent_k(k)/W.dot_centrifugal', 520, 55)])
S.row([('defq', 'η.centrifugal := ' + vcall('η.cent_k', NC), 460, vh(NC))])

# ------------------------------------------------------------------ uncertainty
S.space(10)
S.text('Uncertainty Analysis (Kline-McClintock method)', bold=True, h=26)
S.text('Magnetic pump discharge head, h = P/(ρg): ∂h/∂P = 1/(ρg), ∂h/∂ρ = -h/ρ')
S.row([('fn', 'U.h_mag_discharge_k(k) := sqrt((U.P_mag_discharge/(ρ.water*g))^2+(h.mag_discharge_k(k)/ρ.water*U.ρ_water)^2)', 760, 60)])
S.row([('defq', 'U.h_mag_discharge := ' + vcall('U.h_mag_discharge_k', NM), 560, vh(NM), "'m")])
S.text('Magnetic pump zenith head, h = Q²/(2gA²) + z with A = πD²/4: ∂h/∂Q = Q/(gA²), '
       '∂h/∂D = -2Q²/(gA²D), ∂h/∂z = 1')
S.row([('fn', 'U.h_mag_zenith_k(k) := sqrt((el(Q.magnetic,k)/(g*A.pipe^2)*U.Q_magnetic)^2+(2*el(Q.magnetic,k)^2/(g*A.pipe^2*D.pipe)*U.D_pipe)^2+U.z_zenith^2)', 900, 60)])
S.row([('defq', 'U.h_mag_zenith := ' + vcall('U.h_mag_zenith_k', NM), 540, vh(NM), "'m")])
S.text('Magnetic pump efficiency, η = ρQ(V²/2 + gz)/Ẇ: ∂η/∂ρ = η/ρ, ∂η/∂Ẇ = -η/Ẇ, ∂η/∂Q = ρ(e + V²)/Ẇ, '
       '∂η/∂z = ρQg/Ẇ, ∂η/∂D = -2ρQV²/(Ẇ D)')
S.row([('fn', 'U.η_mag_k(k) := sqrt((η.mag_k(k)/ρ.water*U.ρ_water)^2+(η.mag_k(k)/el(W.magnetic,k)*el(U.W_magnetic,k))^2'
       '+(ρ.water*(e.mag_zenith_k(k)+V.zenith_k(k)^2)/el(W.magnetic,k)*U.Q_magnetic)^2'
       '+(ρ.water*el(Q.magnetic,k)*g/el(W.magnetic,k)*U.z_zenith)^2'
       '+(2*ρ.water*el(Q.magnetic,k)*V.zenith_k(k)^2/(el(W.magnetic,k)*D.pipe)*U.D_pipe)^2)', 1300, 60)])
S.row([('defq', 'U.η_magnetic := ' + vcall('U.η_mag_k', NM), 480, vh(NM))])
S.text('Centrifugal pump head, h = (P_discharge - P_intake)/(ρg): ∂h/∂P_discharge = 1/(ρg), '
       '∂h/∂P_intake = -1/(ρg), ∂h/∂ρ = -h/ρ')
S.row([('fn', 'U.h_cent_k(k) := sqrt((U.P_cent_discharge/(ρ.water*g))^2+(U.P_cent_intake/(ρ.water*g))^2+(h.cent_k(k)/ρ.water*U.ρ_water)^2)', 860, 60)])
S.row([('defq', 'U.h_centrifugal := ' + vcall('U.h_cent_k', NC), 540, vh(NC), "'m")])
S.text('Centrifugal pump power and efficiency. Since ṁ·e = ρQ·ΔP/ρ = Q·ΔP, the density cancels and '
       'η = QΔP/(IV): ∂η/∂Q = ΔP/Ẇ, ∂η/∂P_discharge = Q/Ẇ, ∂η/∂P_intake = -Q/Ẇ, ∂η/∂Ẇ = -η/Ẇ')
S.row([('defq', 'U.W_dot_centrifugal := sqrt((V.centrifugal*U.I_centrifugal)^2+(I.centrifugal*U.V_centrifugal)^2)', 640, 45, "'W")])
S.row([('fn', 'U.η_cent_k(k) := sqrt((ΔP.cent_k(k)/W.dot_centrifugal*el(U.Q_centrifugal,k))^2'
       '+(el(Q.centrifugal,k)/W.dot_centrifugal*U.P_cent_discharge)^2'
       '+(el(Q.centrifugal,k)/W.dot_centrifugal*U.P_cent_intake)^2'
       '+(η.cent_k(k)/W.dot_centrifugal*U.W_dot_centrifugal)^2)', 1100, 60)])
S.row([('defq', 'U.η_centrifugal := ' + vcall('U.η_cent_k', NC), 500, vh(NC))])

# ------------------------------------------------------------------ results in SI for plotting
E = S.env.vars
g_gpm = UNITS['gal'] / 60
qm, qc = np.array(Q_mag), np.array(Q_c)
hmd, uhmd = E['h.mag_discharge'], E['U.h_mag_discharge']
hmz, uhmz = E['h.mag_zenith'], E['U.h_mag_zenith']
etam, uetam = E['η.magnetic'] * 100, E['U.η_magnetic'] * 100
hc, uhc = E['h.centrifugal'], E['U.h_centrifugal']
etac, uetac = E['η.centrifugal'] * 100, E['U.η_centrifugal'] * 100
qs, hs, uhs = np.array(Q_sys), np.array(h_sys), np.array(U_h_sys)


def crossing(qa, ha, qb, hb):
    """Flowrate where two piecewise-linear curves cross, within their overlap."""
    lo, hi = max(qa.min(), qb.min()), min(qa.max(), qb.max())
    if lo >= hi:
        return None
    grid = np.linspace(lo, hi, 4001)
    oa, ob = np.argsort(qa), np.argsort(qb)
    d = np.interp(grid, qa[oa], ha[oa]) - np.interp(grid, qb[ob], hb[ob])
    idx = np.where(np.sign(d[:-1]) != np.sign(d[1:]))[0]
    if len(idx) == 0:
        return None
    i = idx[0]
    q = grid[i] - d[i] * (grid[i + 1] - grid[i]) / (d[i + 1] - d[i])
    return q, float(np.interp(q, qb[ob], hb[ob]))


op_mag = crossing(qm, hmz, qs, hs)
op_cent = crossing(qc, hc, qs, hs)

# ------------------------------------------------------------------ excel workbook (source of the plots)
from openpyxl import Workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.error_bar import ErrorBars
from openpyxl.chart.data_source import NumDataSource, NumRef
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties

DATA = 'Data'
COLORS = ['1F77B4', 'FF7F0E', '2CA02C', 'D62728']
SYMBOLS = ['circle', 'square', 'triangle', 'diamond']
# table layout on the Data sheet: (first data row, last data row)
MAG = (3, 2 + NM)
CEN = (MAG[1] + 4, MAG[1] + 3 + NC)
SYS = (CEN[1] + 4, CEN[1] + 3 + len(Q_sys))


def fill_data(ws):
    def table(title, row0, hdr, rows):
        ws.cell(row0 - 2, 1, title)
        for j, h in enumerate(hdr):
            ws.cell(row0 - 1, j + 1, h)
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                ws.cell(row0 + i, j + 1, v if isinstance(v, str) else round(float(v), 5))
    table('Magnetic Drive Pump', MAG[0],
          ['Pipe', 'Q_magnetic (GPM)', 'h_discharge (m)', 'U h_discharge (m)', 'h_zenith (m)',
           'U h_zenith (m)', 'η_magnetic (%)', 'U η_magnetic (%)'],
          [[f'Pipe {k + 1}', qm[k], hmd[k], uhmd[k], hmz[k], uhmz[k], etam[k], uetam[k]] for k in range(NM)])
    table('Transparent Centrifugal Pump', CEN[0],
          ['Setting', 'Q_centrifugal (GPM)', 'h_centrifugal (m)', 'U h_centrifugal (m)',
           'η_centrifugal (%)', 'U η_centrifugal (%)'],
          [[f'{pct[k]}%', qc[k], hc[k], uhc[k], etac[k], uetac[k]] for k in range(NC)])
    table('System Curve (Lab 6, Part B)', SYS[0],
          ['Test', 'Q_system (GPM)', 'h_L,system (m)', 'U h_L,system (m)'],
          [[f'Test {k + 1}', qs[k], hs[k], uhs[k]] for k in range(len(Q_sys))])


def chart(ws, title, xt, yt, xr, yr, series):
    """series: list of (name, rows, xcol, ycol, ucol). Axis ranges (lo, hi, step)."""
    ch = ScatterChart(); ch.title = title; ch.style = 2; ch.scatterStyle = 'lineMarker'
    ch.x_axis.title, ch.y_axis.title = xt, yt
    ch.x_axis.scaling.min, ch.x_axis.scaling.max, ch.x_axis.majorUnit = xr
    ch.y_axis.scaling.min, ch.y_axis.scaling.max, ch.y_axis.majorUnit = yr
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ch.legend.position = 'b'
    ch.height, ch.width = 11, 18
    for j, (name, (r0, r1), xc, yc, uc) in enumerate(series):
        s = Series(Reference(ws, min_col=yc, min_row=r0, max_row=r1),
                   Reference(ws, min_col=xc, min_row=r0, max_row=r1), title=name)
        s.marker.symbol = SYMBOLS[j]; s.marker.size = 7
        s.marker.graphicalProperties = GraphicalProperties(solidFill=COLORS[j])
        s.marker.graphicalProperties.line.solidFill = COLORS[j]
        s.graphicalProperties.line.noFill = True
        L = chr(64 + uc)
        ref = NumRef(f=f"'{DATA}'!${L}${r0}:${L}${r1}")
        s.errBars = ErrorBars(errDir='y', errBarType='both', errValType='cust', noEndCap=False,
                              plus=NumDataSource(numRef=ref), minus=NumDataSource(numRef=ref),
                              spPr=GraphicalProperties(ln=LineProperties(solidFill=COLORS[j], w=12700)))
        ch.series.append(s)
    return ch


def charts(ws):
    Qx = 'Flowrate, Q (GPM)'
    return [
        ('magnetic_heads', chart(ws, 'Magnetic Drive Pump Curves', Qx, 'Pump Head, h (m)', (0.5, 1.8, 0.1), (0, 3.5, 0.5), [
            ('Pump Head at Discharge', MAG, 2, 3, 4), ('Pump Head at Pipe Zenith', MAG, 2, 5, 6)])),
        ('magnetic_efficiency', chart(ws, 'Magnetic Drive Pump Efficiency', Qx, 'Efficiency, η (%)', (0.5, 1.8, 0.1), (0, 6, 1), [
            ('Magnetic Drive Pump Efficiency', MAG, 2, 7, 8)])),
        ('centrifugal_efficiency', chart(ws, 'Centrifugal Pump Efficiency', Qx, 'Efficiency, η (%)', (0, 50, 5), (0, 25, 5), [
            ('Centrifugal Pump Efficiency', CEN, 2, 5, 6)])),
        ('pump_vs_system', chart(ws, 'Pump Curves vs. Lab 6 System Curve', Qx, 'Head, h (m)', (0, 50, 5), (0, 30, 5), [
            ('Magnetic Drive Pump Curve', MAG, 2, 5, 6), ('Centrifugal Pump Curve', CEN, 2, 3, 4),
            ('System Curve', SYS, 2, 3, 4)])),
        ('pump_vs_system_zoom', chart(ws, 'Pump Curves vs. System Curve (0 to 10 GPM detail)', Qx, 'Head, h (m)', (0, 10, 1), (0, 30, 5), [
            ('Magnetic Drive Pump Curve', MAG, 2, 5, 6), ('Centrifugal Pump Curve', CEN, 2, 3, 4),
            ('System Curve', SYS, 2, 3, 4)])),
    ]


xlsx = os.path.join(OUT, 'Lab_08_-_Pump_Curve_Plots.xlsx')
wb = Workbook(); ws = wb.active; ws.title = DATA; fill_data(ws)
cs = wb.create_sheet('Charts')
for i, (_, ch) in enumerate(charts(ws)):
    cs.add_chart(ch, f'A{1 + 24 * i}')
wb.save(xlsx)

# Render each chart to PNG with LibreOffice: one chart per landscape sheet -> PDF page -> PNG.
import pymupdf
from PIL import Image, ImageChops
tmp = tempfile.mkdtemp()
wb = Workbook(); ws = wb.active; ws.title = DATA; fill_data(ws)
names = []
for i, (name, ch) in enumerate(charts(ws)):
    sh = wb.create_sheet(name, i); sh.add_chart(ch, 'A1'); names.append(name)
    sh.page_setup.orientation = 'landscape'
    sh.page_margins.left = sh.page_margins.right = sh.page_margins.top = sh.page_margins.bottom = 0.3
wb.active = 0
wb.save(os.path.join(tmp, 'charts.xlsx'))
subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp,
                os.path.join(tmp, 'charts.xlsx')], check=True, capture_output=True)
doc = pymupdf.open(os.path.join(tmp, 'charts.pdf'))
pngs = []
for i, name in enumerate(names):
    path = os.path.join(OUT, f'Lab_08_-_{name}.png')
    doc[i].get_pixmap(dpi=110).save(path)
    im = Image.open(path).convert('RGB')
    bbox = ImageChops.difference(im, Image.new('RGB', im.size, 'white')).getbbox()
    im = im.crop((max(bbox[0] - 12, 0), max(bbox[1] - 12, 0), bbox[2] + 12, bbox[3] + 12))
    im.thumbnail((720, 720)); im.save(path)
    pngs.append((path, im.size))
shutil.rmtree(tmp)

S.space(10)
S.text('Plots', bold=True, h=26)
S.text('Charts generated in MS Excel (Lab_08_-_Pump_Curve_Plots.xlsx) and pasted here. Heads are in meters, '
       'flowrates in gallons per minute, and the error bars show ± uncertainty.')
captions = ['Magnetic drive pump curves: head at the discharge port and at the pipe zenith',
            'Magnetic drive pump efficiency vs. flowrate',
            'Centrifugal pump efficiency vs. flowrate',
            'Magnetic and centrifugal pump curves overlaid on the Lab 6 system curve',
            'Same comparison, zoomed to 0 to 10 GPM to show the operating points']
for (path, (w, h)), cap in zip(pngs, captions):
    S.text(cap)
    S.picture(path, w, h)

# ------------------------------------------------------------------ Q&A
max_mag_head = hmz.max()
qa = [
    'Why the discharge head differs from the zenith head: between the pressure gauge at the pump discharge '
    'and the top of the pipe, the water loses energy to friction along the vertical pipe and to the minor '
    'losses in the tee, rotameter, ball valve, elbows, and the "Y" fitting at the top. The discharge head is '
    'the energy the pump delivers, while the zenith head is what is left after those losses, so the discharge '
    'head should be the larger of the two. The difference also grows with height because a taller pipe has '
    'more length for friction. In our data the zenith head is about equal to the discharge head for pipe 5 '
    'and larger than it for pipe 6, which is not physically possible, so part of the difference there is '
    'measurement error, for example the gauge not sitting at the same height as the discharge port, the '
    'zenith height being measured to the wrong reference point, or the gauge tare drifting.',
    'Factors that contribute to pump efficiency: (1) hydraulic losses - friction and turbulence as the water '
    'flows through the impeller and casing, plus shock losses when the flow enters the blades at the wrong '
    'angle; (2) volumetric (leakage) losses - water that recirculates through the clearance between the '
    'impeller and casing instead of leaving the discharge; (3) mechanical losses - bearing, seal, and disk '
    'friction on the spinning impeller; and (4) the electric motor and drive losses, since we measured '
    'electrical power rather than shaft power. Impeller design, running clearances, rotational speed, '
    'fluid viscosity, and pump size all affect these losses. A magnetic drive pump also loses energy in the '
    'magnetic coupling. (Reference: Cengel and Cimbala, Fluid Mechanics: Fundamentals and Applications, '
    'Chapter 14, Turbomachinery.)',
    'Why the efficiency changes with flowrate: efficiency is the useful power ρgQh divided by the input '
    'power. At zero flow (valve closed) the pump still draws power but delivers no flow, so the efficiency '
    'is zero. As the flow increases the useful power rises faster than the losses, so the efficiency rises '
    'to a maximum at the best efficiency point (BEP), where the flow enters the impeller blades at their '
    'design angle. Beyond the BEP the head falls off and the friction and shock losses grow with the square '
    'of the velocity, so the efficiency drops again. For the centrifugal testbench the input power was the '
    'same at every setting (237.9 V × 4.1 A), so its efficiency followed Q·ΔP and was highest at the largest '
    'flowrate we could reach. The magnetic pump efficiency peaked at about 5.6% near 1.05 to 1.15 GPM '
    '(pipes 4 and 5) and fell off at both higher and lower flowrates, which shows its best efficiency point.',
]
if op_mag:
    qa.append(f'Magnetic drive pump vs. the Lab 6 system: yes, the magnetic pump curve crosses the system curve, '
              f'so the pump could push flow through the Lab 6 piping. The operating point is at about '
              f'{op_mag[0]:.2f} GPM with a head of about {op_mag[1]:.2f} m. This is at the very low-flow end '
              f'of the system curve, well below the 3.5 GPM maximum we measured in Lab 6, because the '
              f'magnetic pump can only produce about {max_mag_head:.1f} m of head.')
else:
    qa.append(f'Magnetic drive pump vs. the Lab 6 system: the magnetic pump curve does not cross the system curve '
              f'within the measured flowrates. The magnetic pump only produces about {max_mag_head:.1f} m of head.')
if op_cent:
    qa.append(f'Centrifugal pump vs. the Lab 6 system: yes, the centrifugal pump curve crosses the system curve, so '
              f'it could push flow through the Lab 6 piping. The operating point is at about {op_cent[0]:.2f} GPM '
              f'with a head of about {op_cent[1]:.1f} m. Because we have no centrifugal readings between 0 and '
              f'8 GPM, this point uses a straight line between the 0% and 20% settings, so it is an estimate.')
else:
    qa.append('Centrifugal pump vs. the Lab 6 system: the centrifugal pump curve does not cross the system curve '
              'within the measured flowrates.')
S.space(10)
S.text('Q/A', bold=True, h=26)
for q in qa:
    S.text(q, w=760)

S.save(os.path.join(OUT, 'Lab_08_-_Pump_Curve.sm'), 'Grant Bellon', '9d3f6a20-5c7e-4b81-a2f4-6e1b8c03d957')

np.set_printoptions(precision=4, suppress=True)
print('rho', E['ρ.water'], '+-', E['U.ρ_water'])
print('h_md', hmd, uhmd); print('h_zen', hmz, uhmz); print('eta_mag %', etam, uetam)
print('h_c', hc, uhc); print('eta_c %', etac, uetac); print('W_c', E['W.dot_centrifugal'])
print('op_mag', op_mag, 'op_cent', op_cent)
